"""Text extraction and chunking, done entirely on in-memory bytes."""
from __future__ import annotations

import html
import io
import re

SUPPORTED = {".pdf", ".docx", ".xlsx", ".pptx", ".txt", ".md", ".markdown", ".csv", ".json", ".log", ".html", ".htm",
             ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
IMAGES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}


class IngestError(ValueError):
    pass


def extension(name: str) -> str:
    m = re.search(r"(\.[A-Za-z0-9]+)$", name or "")
    return m.group(1).lower() if m else ""


def _decode(data: bytes) -> str:
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise IngestError("Could not decode text file.")


def _pdf(data: bytes) -> str:
    try:
        import pymupdf
    except ImportError:
        pymupdf = None
    if pymupdf is not None:
        try:
            doc = pymupdf.open(stream=data, filetype="pdf")
        except Exception as exc:  # noqa: BLE001
            raise IngestError("Could not read PDF.") from exc
        if doc.needs_pass:
            raise IngestError("PDF is password-protected.")
        text = "\n\n".join(page.get_text("text", sort=True) for page in doc).strip()
        doc.close()
        if not text:
            raise IngestError("PDF has no extractable text (it may be a scan).")
        return text
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:  # noqa: BLE001
                raise IngestError("PDF is password-protected.") from exc
        pages = [(p.extract_text() or "") for p in reader.pages]
    except IngestError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise IngestError("Could not read PDF.") from exc
    text = "\n\n".join(pages).strip()
    if not text:
        raise IngestError("PDF has no extractable text (it may be a scan).")
    return text


def _docx(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise IngestError("Could not read DOCX.") from exc
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n\n".join(parts)


def _html(data: bytes) -> str:
    raw = _decode(data)
    raw = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", raw)
    raw = re.sub(r"(?s)<[^>]+>", " ", raw)
    return html.unescape(raw)


def extract_text(name: str, data: bytes) -> str:
    ext = extension(name)
    if ext not in SUPPORTED:
        raise IngestError(f"Unsupported file type '{ext or name}'.")
    if ext == ".pdf":
        text = _pdf(data)
    elif ext == ".xlsx":
        from .office import xlsx_text
        try:
            text = xlsx_text(data)
        except Exception as exc:  # noqa: BLE001
            raise IngestError("Could not read this Excel file.") from exc
    elif ext == ".pptx":
        from .office import pptx_text
        try:
            text = pptx_text(data)
        except Exception as exc:  # noqa: BLE001
            raise IngestError("Could not read this PowerPoint file.") from exc
    elif ext == ".docx":
        text = _docx(data)
    elif ext in {".html", ".htm"}:
        text = _html(data)
    else:
        text = _decode(data)
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise IngestError("The file contains no text.")
    return text


def chunk_text(text: str, size: int = 1000, overlap: int = 150) -> list[str]:
    """Split on paragraph boundaries, then pack into ~size-char chunks with overlap."""
    overlap = max(0, min(overlap, size // 2))
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    pieces: list[str] = []
    for p in paras:
        if len(p) <= size:
            pieces.append(p)
            continue
        # Long paragraph: split on sentences, then hard-split if still too long.
        sentences = re.split(r"(?<=[.!?])\s+", p)
        for s in sentences:
            while len(s) > size:
                pieces.append(s[:size])
                s = s[size - overlap:]
            if s:
                pieces.append(s)

    chunks: list[str] = []
    current = ""
    for piece in pieces:
        if current and len(current) + len(piece) + 2 > size:
            chunks.append(current)
            tail = current[-overlap:] if overlap else ""
            current = (tail + "\n\n" + piece).strip() if tail else piece
            if len(current) > size:  # tail + piece too long; drop the tail
                current = piece
        else:
            current = f"{current}\n\n{piece}" if current else piece
    if current:
        chunks.append(current)
    return chunks
