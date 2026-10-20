"""Turn an answer into a PDF or Word file, entirely in memory.

Nothing is written to disk: the file is built in a BytesIO buffer and sent
straight back to the browser.
"""
from __future__ import annotations

import io
import os
import re

FONT_DIRS = ["/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/dejavu"]

_CITE = re.compile(r"\s?\[\d{1,2}\]")
_PREAMBLE = re.compile(
    r"^(here is|here's|hier ist|hier sind|voici|sure[,!]|gerne[,!]|okay[,!]|natürlich[,!]).{0,120}:\s*$",
    re.IGNORECASE,
)


def clean_text(answer: str) -> str:
    """Keep only the document part of an answer.

    The prompt asks the model to put its "what I changed" note after a line
    with only `---`, so everything after the last such line is dropped. A short
    intro line ("Hier ist die überarbeitete Version:") is dropped too, as are
    citation markers and Markdown emphasis.
    """
    text = answer.replace("\r\n", "\n").strip()
    parts = re.split(r"\n\s*-{3,}\s*\n", text)
    if len(parts) > 1:
        text = "\n---\n".join(parts[:-1]).strip()
    lines = text.split("\n")
    if lines and _PREAMBLE.match(lines[0].strip()):
        lines = lines[1:]
    text = "\n".join(lines).strip()
    text = _CITE.sub("", text)
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"(?m)^#{1,6}\s+", "", text)
    text = text.replace("```", "")
    return text.strip() or answer.strip()


def _paragraphs(text: str) -> list[str]:
    return [p.strip("\n") for p in re.split(r"\n\s*\n", text)]


def _font_path(name: str) -> str | None:
    for d in FONT_DIRS:
        p = os.path.join(d, name)
        if os.path.exists(p):
            return p
    return None


def to_pdf(text: str, title: str = "Paradox") -> bytes:
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_margins(25, 25, 25)
    pdf.set_auto_page_break(True, margin=22)
    pdf.set_title(title)
    pdf.set_creator("Paradox")
    pdf.add_page()

    regular = _font_path("DejaVuSans.ttf")
    if regular:
        pdf.add_font("Body", "", regular)
        family = "Body"
    else:  # fall back to a core font; it only covers Latin-1 characters
        family = "Helvetica"
        text = text.translate({0x2013: "-", 0x2014: "-", 0x201E: '"', 0x201C: '"', 0x201D: '"',
                               0x2018: "'", 0x2019: "'", 0x2022: "-", 0x20AC: "EUR"})
        text = text.encode("latin-1", "replace").decode("latin-1")
    pdf.set_font(family, size=11)

    width = pdf.w - pdf.l_margin - pdf.r_margin
    for para in _paragraphs(text):
        if not para.strip():
            continue
        for line in para.split("\n"):
            pdf.multi_cell(width, 5.6, line if line.strip() else " ", align="L", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3.2)
    return bytes(pdf.output())


def to_docx(text: str) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(8)
    for para in _paragraphs(text):
        if not para.strip():
            continue
        p = doc.add_paragraph()
        lines = para.split("\n")
        for i, line in enumerate(lines):
            run = p.add_run(line)
            if i < len(lines) - 1:
                run.add_break()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
